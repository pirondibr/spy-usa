# -*- coding: utf-8 -*-
"""Spy USA pipeline: briefing → (concorrentes?) → SEO → Brand → Google Ads → Meta → Social."""

from __future__ import annotations

import os
import subprocess
import sys
import threading
import time
from pathlib import Path
from typing import Any, Callable, Optional

from parse_input import ParsedInput, normalize_url, slugify_client
from report_builder import build_early_briefing_competitors, build_report_from_xlsx
from report_html import render_report_html

APP_DIR = Path(__file__).resolve().parent
FINAL_DIR = APP_DIR / "scripts"
SCRIPT_ENTENDER = FINAL_DIR / "1- entender o cliente.py"
SCRIPT_CONCORRENTES = FINAL_DIR / "3 - concorrentes Geral.py"
SCRIPT_SEO = FINAL_DIR / "5b - seo organico.py"
SCRIPT_BRAND = FINAL_DIR / "5c - brand search.py"
SCRIPT_GOOGLE = FINAL_DIR / "5a - google ads.py"
SCRIPT_META = FINAL_DIR / "5d - meta ads.py"
SCRIPT_LINKEDIN = FINAL_DIR / "5e - linkedin ads.py"
SCRIPT_BRAND_DFS = FINAL_DIR / "5g - brand dataforseo.py"
SCRIPT_SIMILARWEB = FINAL_DIR / "5h - similarweb.py"
SCRIPT_SOCIAL = FINAL_DIR / "5f - social ig yt.py"

_DATA_DIR = (os.environ.get("RADAR_DATA_DIR") or os.environ.get("SPY_DATA_DIR") or "").strip()
_OUT_ROOT = (Path(_DATA_DIR) / "outputs") if _DATA_DIR else (APP_DIR / "outputs")
METRICAS_DIR = _OUT_ROOT / "metricas"
BRIEFING_DIR = _OUT_ROOT / "entender"
CONCORRENTES_DIR = _OUT_ROOT / "concorrentes"
RUNS_DIR = _OUT_ROOT / "runs"

EmitFn = Callable[..., None]

STEP_DEFS = [
    {"id": "briefing_concorrentes", "label": "Briefing (+ competitors)", "index": 1, "eta_live": 180},
    {"id": "seo", "label": "SEO organic (US)", "index": 2, "eta_live": 30},
    {"id": "brand", "label": "Brand search Semrush (US)", "index": 3, "eta_live": 30},
    {"id": "brand_dfs", "label": "Brand DataForSEO (US)", "index": 4, "eta_live": 40},
    {"id": "similarweb", "label": "SimilarWeb Traffic", "index": 5, "eta_live": 60},
    {"id": "google_ads", "label": "Google Ads Library (US)", "index": 6, "eta_live": 180},
    {"id": "meta", "label": "Meta Ads (US)", "index": 7, "eta_live": 180},
    {"id": "linkedin", "label": "LinkedIn Ads Library (US)", "index": 8, "eta_live": 240},
    {"id": "instagram", "label": "Instagram", "index": 9, "eta_live": 120},
    {"id": "youtube", "label": "YouTube", "index": 10, "eta_live": 20},
    {"id": "tiktok", "label": "TikTok", "index": 11, "eta_live": 40},
]
TOTAL_STEPS = len(STEP_DEFS)


def _load_brand_dfs(slug: str) -> Optional[dict]:
    path = METRICAS_DIR / slug / f"brand-dataforseo-{slug}.json"
    if not path.exists():
        return None
    try:
        import json

        return json.loads(path.read_text(encoding="utf-8"))
    except Exception:
        return None


def _attach_brand_dfs(report: dict, slug: str) -> dict:
    data = _load_brand_dfs(slug)
    if not data:
        return report
    rows_raw = data.get("rows") or []
    from urllib.parse import quote

    rows = []
    for r in rows_raw:
        vol = r.get("latest_volume")
        g1 = r.get("growth_1y_pct")
        kw = (r.get("brand_keyword") or "").strip()
        kw_url = (
            f"https://www.semrush.com/analytics/keywordoverview/?q={quote(kw)}&db=us"
            if kw
            else ""
        )
        rows.append({
            "name": r.get("name") or r.get("brand_keyword"),
            "domain": r.get("domain"),
            "is_client": bool(r.get("is_client")),
            "similaridade": r.get("similaridade") or ("Cliente" if r.get("is_client") else "—"),
            "brand_keyword": r.get("brand_keyword"),
            "traffic": vol or 0,
            "traffic_fmt": f"{vol:,}" if isinstance(vol, int) else (str(vol) if vol is not None else "—"),
            "growth_fmt": (f"{g1:+.1f}%" if isinstance(g1, (int, float)) else "n/a"),
            "latest_label": r.get("latest_label"),
            "volume_1y": r.get("volume_1y"),
            "volume_2y": r.get("volume_2y"),
            "volume_3y": r.get("volume_3y"),
            "volume_5y": r.get("volume_5y"),
            "growth_1y_pct": r.get("growth_1y_pct"),
            "growth_2y_pct": r.get("growth_2y_pct"),
            "growth_3y_pct": r.get("growth_3y_pct"),
            "growth_5y_pct": r.get("growth_5y_pct"),
            "site_title": r.get("site_title"),
            "url": kw_url,
        })
    client = next((r for r in rows_raw if r.get("is_client")), rows_raw[0] if rows_raw else {})
    report["brand_dataforseo"] = {
        "source": "dataforseo",
        "market": "US",
        "latest_label": data.get("latest_label") or client.get("latest_label"),
        "client_keyword": data.get("client_keyword") or client.get("brand_keyword"),
        "leader": (rows[0].get("name") if rows else "—"),
        "total_traffic_fmt": rows[0].get("traffic_fmt") if rows else "—",
        "unit": "monthly searches",
        "rows": rows,
        "insight": (
            f"Brand keyword «{data.get('client_keyword') or client.get('brand_keyword') or '—'}» "
            f"· latest month {data.get('latest_label') or client.get('latest_label') or 'n/a'} "
            f"(DataForSEO Google Ads US). Growth vs same month 1y/2y/3y/5y."
        ),
        "analysis_title": "Brand Search — DataForSEO",
        "raw": data,
    }
    return report


def _load_similarweb(slug: str) -> Optional[dict]:
    path = METRICAS_DIR / slug / f"similarweb-{slug}.json"
    if not path.exists():
        return None
    try:
        import json

        return json.loads(path.read_text(encoding="utf-8"))
    except Exception:
        return None


def _attach_similarweb(report: dict, slug: str) -> dict:
    data = _load_similarweb(slug)
    if not data:
        return report
    rows_raw = data.get("rows") or []
    rows: list[dict] = []
    for r in rows_raw:
        rows.append({
            "name": r.get("name") or r.get("domain"),
            "domain": r.get("domain"),
            "is_client": bool(r.get("is_client")),
            "similaridade": r.get("similaridade") or ("Cliente" if r.get("is_client") else "—"),
            "traffic": r.get("traffic") or r.get("latest_visits") or 0,
            "traffic_fmt": r.get("traffic_fmt") or "—",
            "snapshot_date": r.get("snapshot_date"),
            "snapshot_label": r.get("snapshot_label") or "—",
            "bounce_fmt": r.get("bounce_fmt") or "—",
            "pages_fmt": r.get("pages_fmt") or "—",
            "time_fmt": r.get("time_fmt") or "—",
            "is_data_from_ga": bool(r.get("is_data_from_ga")),
            "visits_monthly": r.get("visits_monthly") or [],
            "sources": r.get("sources") or [],
            "countries": r.get("countries") or [],
            "error": r.get("error"),
            "url": r.get("url") or (f"https://{r.get('domain')}/" if r.get("domain") else ""),
        })
    client = next((r for r in rows if r.get("is_client")), rows[0] if rows else {})
    ok_rows = [r for r in rows if not r.get("error")]
    leader = ok_rows[0] if ok_rows else client
    report["similarweb"] = {
        "source": "similarweb_rapidapi",
        "market": "global",
        "latest_label": data.get("latest_label") or client.get("snapshot_label"),
        "client_domain": data.get("client_domain") or client.get("domain"),
        "leader": leader.get("name") if leader else "—",
        "total_traffic_fmt": (client.get("traffic_fmt") if client else "—"),
        "unit": "visits/mo",
        "rows": rows,
        "insight": (
            f"SimilarWeb traffic for «{client.get('domain') or '—'}» "
            f"· snapshot {client.get('snapshot_label') or data.get('latest_label') or 'n/a'} "
            f"· bounce {client.get('bounce_fmt') or '—'} · "
            f"{client.get('pages_fmt') or '—'} pages/visit · "
            f"{client.get('time_fmt') or '—'} on site."
        ),
        "analysis_title": "SimilarWeb Traffic",
    }
    return report


def find_metricas_xlsx(slug: str) -> Optional[Path]:
    folder = METRICAS_DIR / slug
    if not folder.exists():
        return None
    files = sorted(
        folder.glob(f"metricas-concorrentes-{slug}*.xlsx"),
        key=lambda p: p.stat().st_mtime,
        reverse=True,
    )
    if not files:
        files = sorted(folder.glob("metricas-concorrentes-*.xlsx"), key=lambda p: p.stat().st_mtime, reverse=True)
    return files[0] if files else None


def find_briefing_xlsx(slug: str) -> Optional[Path]:
    folder = BRIEFING_DIR / slug
    if not folder.exists():
        return None
    files = sorted(
        folder.glob(f"briefing-cliente-{slug}*.xlsx"),
        key=lambda p: p.stat().st_mtime,
        reverse=True,
    )
    if not files:
        files = sorted(folder.glob("briefing-cliente-*.xlsx"), key=lambda p: p.stat().st_mtime, reverse=True)
    return files[0] if files else None


def find_concorrentes_xlsx(slug: str) -> Optional[Path]:
    folder = CONCORRENTES_DIR / slug
    if not folder.exists():
        return None
    nacional = sorted(
        folder.glob(f"concorrentes-all-{slug}-nacional*.xlsx"),
        key=lambda p: p.stat().st_mtime,
        reverse=True,
    )
    if nacional:
        return nacional[0]
    files = sorted(
        folder.glob(f"concorrentes-all-{slug}*.xlsx"),
        key=lambda p: p.stat().st_mtime,
        reverse=True,
    )
    if not files:
        files = sorted(folder.glob("concorrentes-all-*.xlsx"), key=lambda p: p.stat().st_mtime, reverse=True)
    return files[0] if files else None


def clear_slug_cache(slug: str) -> dict[str, Any]:
    import re
    import shutil

    safe = re.sub(r"[^a-z0-9_-]+", "", (slug or "").strip().lower())[:80]
    if not safe:
        raise ValueError("Invalid slug")
    deleted: list[str] = []
    for base in (METRICAS_DIR, BRIEFING_DIR, CONCORRENTES_DIR):
        folder = (base / safe).resolve()
        try:
            folder.relative_to(base.resolve())
        except ValueError as e:
            raise ValueError(f"Path outside outputs: {folder}") from e
        if folder.exists() and folder.is_dir():
            shutil.rmtree(folder)
            deleted.append(str(folder))
    return {"slug": safe, "deleted": deleted, "cleared": bool(deleted)}


def build_solo_national_xlsx(slug: str, client_url: str, client_name: str) -> Path:
    """Minimal concorrentes XLSX with only the client (Semrush US traffic, no competitor discovery)."""
    sys.path.insert(0, str(FINAL_DIR))
    sys.path.insert(0, str(FINAL_DIR / "vendor"))
    from workspace_paths import setup_workspace  # noqa: E402

    setup_workspace()
    from find_concorrentes import build_3period_traffic  # noqa: E402
    from find_concorrentes_all import _write_growth_sheet, _write_traffic_sheet  # noqa: E402
    from find_concorrentes import read_briefing_from_xlsx  # noqa: E402
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from datetime import datetime

    briefing_path = find_briefing_xlsx(slug)
    if not briefing_path:
        raise RuntimeError(f"Briefing not found for solo mode: {slug}")
    briefing = read_briefing_from_xlsx(briefing_path)
    url = briefing.get("url") or client_url
    profile = briefing.get("perfil") or "—"

    emit_log = f"[SOLO] Semrush traffic for {url} only (no competitors)"
    print(emit_log, flush=True)
    traffic = build_3period_traffic(url, [], profile)

    out_dir = CONCORRENTES_DIR / slug
    out_dir.mkdir(parents=True, exist_ok=True)
    out_path = out_dir / f"concorrentes-all-{slug}-nacional.xlsx"

    wb = Workbook()
    ws = wb.active
    ws.title = "Resumo"
    ws.append(["Spy USA — site only (no competitors)"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([])
    ws.append(["Cliente", client_name])
    ws.append(["URL", url])
    ws.append(["Slug", slug])
    ws.append(["Mode", "site_only"])
    ws.append(["Market", "US"])
    ws.append(["Data", datetime.now().strftime("%d/%m/%Y %H:%M")])

    ws_c = wb.create_sheet("Concorrentes Unificado")
    ws_c.append(["Dominio", "URL", "Titulo", "Similaridade", "Perfil", "Fonte", "Nicho (LLM)"])
    # empty competitors — metrics script allows SPY_SITE_ONLY

    _write_traffic_sheet(wb, traffic)
    _write_growth_sheet(wb, traffic)
    wb.save(out_path)
    print(f"[SOLO] Saved {out_path}", flush=True)
    return out_path


def _run_script(
    cmd: list[str],
    cwd: Path,
    on_log: Optional[Callable[[str], None]] = None,
    timeout_sec: Optional[int] = None,
    env: Optional[dict[str, str]] = None,
) -> None:
    run_env = os.environ.copy()
    if env:
        run_env.update(env)
    proc = subprocess.Popen(
        cmd,
        cwd=str(cwd),
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
        text=True,
        encoding="utf-8",
        errors="replace",
        bufsize=1,
        env=run_env,
    )
    assert proc.stdout is not None
    tail: list[str] = []

    def _reader() -> None:
        for line in proc.stdout:
            line = line.rstrip("\n\r")
            if not line:
                continue
            tail.append(line)
            if len(tail) > 40:
                del tail[:-40]
            if on_log:
                on_log(line)

    reader = threading.Thread(target=_reader, daemon=True)
    reader.start()
    reader.join(timeout=timeout_sec)
    if reader.is_alive():
        proc.kill()
        try:
            proc.wait(timeout=10)
        except Exception:
            pass
        hint = ("\n--- last lines ---\n" + "\n".join(tail[-15:])) if tail else ""
        raise TimeoutError(f"Command exceeded {timeout_sec}s: {' '.join(cmd)}{hint}")
    proc.wait()
    if proc.returncode != 0:
        hint = ("\n--- last lines ---\n" + "\n".join(tail[-20:])) if tail else ""
        raise RuntimeError(f"Command failed ({proc.returncode}): {' '.join(cmd)}{hint}")


def _fmt_eta(seconds: int) -> str:
    if seconds >= 60:
        return f"~{max(1, round(seconds / 60))} min"
    return f"~{seconds}s"


def _emit_progress(
    emit: EmitFn,
    set_step: EmitFn,
    step_id: str,
    state: str,
    detail: str,
) -> None:
    meta = next((s for s in STEP_DEFS if s["id"] == step_id), None)
    index = int(meta["index"]) if meta else 0
    eta = int(meta.get("eta_live") or 60) if meta else 60
    set_step(step_id, state, detail)
    emit(
        "progress",
        step_id=step_id,
        state=state,
        detail=detail,
        index=index,
        total=TOTAL_STEPS,
        label=meta["label"] if meta else step_id,
        eta_seconds=eta,
        eta_label=_fmt_eta(eta),
        progress_label=f"Step {index}/{TOTAL_STEPS}",
    )


def _save_report(
    run_id: str,
    report: dict,
    parsed: ParsedInput,
    *,
    site_only: bool = False,
    include_google_ads: bool = True,
    include_linkedin: bool = True,
    force: bool = False,
) -> Path:
    RUNS_DIR.mkdir(parents=True, exist_ok=True)
    run_dir = RUNS_DIR / run_id
    run_dir.mkdir(parents=True, exist_ok=True)
    import json
    from datetime import datetime, timezone

    created_at = datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
    meta = {
        "id": run_id,
        "slug": parsed.slug,
        "url": parsed.url,
        "company": parsed.company,
        "status": "done",
        "market": "US",
        "site_only": bool(site_only),
        "include_google_ads": bool(include_google_ads),
        "include_linkedin": bool(include_linkedin),
        "force": bool(force),
        "created_at": created_at,
        "html_url": f"/runs/{run_id}/report.html",
    }
    (run_dir / "meta.json").write_text(
        json.dumps(meta, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    (run_dir / "report.json").write_text(
        json.dumps(report, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    html = render_report_html(report, run_meta=meta)
    html_path = run_dir / "report.html"
    html_path.write_text(html, encoding="utf-8")
    return html_path


def run_pipeline(
    parsed: ParsedInput,
    emit: EmitFn,
    set_step: EmitFn,
    run_id: str = "",
    force: bool = False,
    site_only: bool = False,
    include_google_ads: bool = True,
    include_linkedin: bool = True,
) -> dict:
    required = (
        "OPENROUTER_API_KEY",
        "SCRAPINGBEE_API_KEY",
        "DATAFORSEO_USER",
        "DATAFORSEO_PASS",
        "SEMRUSH_API_KEY",
    )
    missing = [k for k in required if not os.environ.get(k, "").strip()]
    if missing:
        raise RuntimeError("Missing API keys: " + ", ".join(missing))

    url = normalize_url(parsed.url) or parsed.url
    if not url:
        raise ValueError("Company URL is required.")
    slug = parsed.slug or slugify_client(url)
    client_name = parsed.company or slug
    script_env = {"SPY_SITE_ONLY": "1" if site_only else "0", "SPY_MARKET": "us"}

    def on_log(line: str) -> None:
        emit("log", line=line)

    if force:
        clear_slug_cache(slug)

    existing = find_metricas_xlsx(slug)
    if existing and existing.exists() and not force:
        emit("log", line=f"[CACHE] Using existing metrics: {existing.name}")
        report = build_report_from_xlsx(
            existing,
            client_name=client_name,
            preferred_competitors=parsed.competitors,
        )
        _emit_progress(emit, set_step, "briefing_concorrentes", "done", "Cached report")
        for sid in ("seo", "brand", "brand_dfs", "similarweb", "google_ads", "meta", "linkedin", "instagram", "youtube", "tiktok"):
            _emit_progress(emit, set_step, sid, "done", "Cached")
        report = _attach_brand_dfs(report, slug)
        report = _attach_similarweb(report, slug)
        if run_id:
            _save_report(
                run_id, report, parsed,
                site_only=site_only,
                include_google_ads=include_google_ads,
                include_linkedin=include_linkedin,
                force=force,
            )
        emit("done", report=report, html_url=f"/runs/{run_id}/report.html" if run_id else "")
        return report

    emit(
        "pipeline_meta",
        mode="live",
        market="US",
        site_only=site_only,
        include_google_ads=include_google_ads,
        include_linkedin=include_linkedin,
        total_steps=TOTAL_STEPS,
        eta_label="~3–8 min" if site_only else "~10–15 min",
    )

    # 1) Briefing + optional competitors
    if site_only:
        _emit_progress(
            emit, set_step, "briefing_concorrentes", "running",
            f"Briefing + modo sem concorrentes: {url}",
        )
    else:
        _emit_progress(
            emit, set_step, "briefing_concorrentes", "running",
            f"Briefing + US competitors: {url}",
        )

    _run_script(
        [sys.executable, str(SCRIPT_ENTENDER), url],
        FINAL_DIR,
        on_log,
        timeout_sec=300,
        env=script_env,
    )
    briefing_xlsx = find_briefing_xlsx(slug)
    if not briefing_xlsx:
        raise RuntimeError(f"Briefing not generated for '{slug}'.")

    if site_only:
        build_solo_national_xlsx(slug, url, client_name)
        early = build_early_briefing_competitors(
            find_briefing_xlsx(slug),
            find_concorrentes_xlsx(slug),
            client_name=client_name,
            preferred_competitors=[],
            fallback_url=url,
        )
        early["competitors_note"] = "Modo sem pesquisar concorrentes — discovery skipped."
        early["competitors_count"] = 0
        emit("partial", section="briefing_concorrentes", data=early, client=client_name)
        _emit_progress(emit, set_step, "briefing_concorrentes", "done", "Sem concorrentes (0)")
    else:
        _run_script(
            [sys.executable, str(SCRIPT_CONCORRENTES), slug, "nacional"],
            FINAL_DIR,
            on_log,
            timeout_sec=600,
            env=script_env,
        )
        early = build_early_briefing_competitors(
            find_briefing_xlsx(slug),
            find_concorrentes_xlsx(slug),
            client_name=client_name,
            preferred_competitors=parsed.competitors,
            fallback_url=url,
        )
        emit("partial", section="briefing_concorrentes", data=early, client=client_name)
        _emit_progress(
            emit, set_step, "briefing_concorrentes", "done",
            f"{early.get('competitors_count', 0)} competitors",
        )

    # 2) SEO
    _emit_progress(emit, set_step, "seo", "running", "Building US SEO ranking...")
    _run_script(
        [sys.executable, str(SCRIPT_SEO), slug],
        FINAL_DIR,
        on_log,
        timeout_sec=180,
        env=script_env,
    )
    xlsx = find_metricas_xlsx(slug)
    if not xlsx:
        raise FileNotFoundError(f"Metrics XLSX not generated for '{slug}'")
    report = build_report_from_xlsx(xlsx, client_name=client_name, preferred_competitors=parsed.competitors)
    emit("partial", section="seo", data=report.get("seo") or {}, client=client_name)
    _emit_progress(emit, set_step, "seo", "done", "SEO ready")

    # 3) Brand
    _emit_progress(emit, set_step, "brand", "running", "Building brand ranking...")
    _run_script(
        [sys.executable, str(SCRIPT_BRAND), slug],
        FINAL_DIR,
        on_log,
        timeout_sec=180,
        env=script_env,
    )
    xlsx = find_metricas_xlsx(slug) or xlsx
    report = build_report_from_xlsx(xlsx, client_name=client_name, preferred_competitors=parsed.competitors)
    emit("partial", section="brand", data=report.get("brand") or {}, client=client_name)
    _emit_progress(emit, set_step, "brand", "done", "Brand Semrush ready")

    # 3b) Brand DataForSEO (keyword from domain+title, monthly history)
    _emit_progress(emit, set_step, "brand_dfs", "running", "Brand DataForSEO US (1y/2y/3y/5y)...")
    try:
        _run_script(
            [sys.executable, str(SCRIPT_BRAND_DFS), slug],
            FINAL_DIR,
            on_log,
            timeout_sec=180,
            env=script_env,
        )
    except Exception as e:
        emit("log", line=f"[BRAND/DFS] Partial failure: {e}")
    report = _attach_brand_dfs(report, slug)
    emit(
        "partial",
        section="brand_dataforseo",
        data=report.get("brand_dataforseo") or {},
        client=client_name,
    )
    _emit_progress(emit, set_step, "brand_dfs", "done", "Brand DataForSEO ready")

    # 3c) SimilarWeb Traffic (RapidAPI)
    _emit_progress(emit, set_step, "similarweb", "running", "SimilarWeb Traffic (RapidAPI)...")
    try:
        _run_script(
            [sys.executable, str(SCRIPT_SIMILARWEB), slug],
            FINAL_DIR,
            on_log,
            timeout_sec=180,
            env=script_env,
        )
    except Exception as e:
        emit("log", line=f"[SIMILARWEB] Partial failure: {e}")
    report = _attach_similarweb(report, slug)
    emit(
        "partial",
        section="similarweb",
        data=report.get("similarweb") or {},
        client=client_name,
    )
    _emit_progress(emit, set_step, "similarweb", "done", "SimilarWeb ready")

    # 4) Google Ads Library (optional)
    if include_google_ads:
        _emit_progress(emit, set_step, "google_ads", "running", "Google Ads Transparency (US)...")
        try:
            _run_script(
                [sys.executable, str(SCRIPT_GOOGLE), slug],
                FINAL_DIR,
                on_log,
                timeout_sec=420,
                env=script_env,
            )
        except Exception as e:
            emit("log", line=f"[GOOGLE] Partial failure: {e}")
        xlsx = find_metricas_xlsx(slug) or xlsx
        report = build_report_from_xlsx(xlsx, client_name=client_name, preferred_competitors=parsed.competitors)
        emit("partial", section="google_ads", data=report.get("google_ads") or {}, client=client_name)
        _emit_progress(emit, set_step, "google_ads", "done", "Google Ads ready")
    else:
        _emit_progress(emit, set_step, "google_ads", "done", "Skipped")

    # 5) Meta
    _emit_progress(emit, set_step, "meta", "running", "Meta Ads Library (US)...")
    try:
        _run_script(
            [sys.executable, str(SCRIPT_META), slug],
            FINAL_DIR,
            on_log,
            timeout_sec=420,
            env=script_env,
        )
    except Exception as e:
        emit("log", line=f"[META] Partial failure: {e}")
    xlsx = find_metricas_xlsx(slug) or xlsx
    report = build_report_from_xlsx(xlsx, client_name=client_name, preferred_competitors=parsed.competitors)
    emit("partial", section="meta", data=report.get("meta") or {}, client=client_name)
    _emit_progress(emit, set_step, "meta", "done", "Meta ready")

    # 6) LinkedIn Ads Library (optional) — empresa → ad → pagante → total
    if include_linkedin:
        _emit_progress(
            emit, set_step, "linkedin", "running",
            "LinkedIn Ads Library (US): find ad → payer → total...",
        )
        try:
            _run_script(
                [sys.executable, str(SCRIPT_LINKEDIN), slug],
                FINAL_DIR,
                on_log,
                timeout_sec=600,
                env=script_env,
            )
        except Exception as e:
            emit("log", line=f"[LINKEDIN] Partial failure: {e}")
        xlsx = find_metricas_xlsx(slug) or xlsx
        report = build_report_from_xlsx(xlsx, client_name=client_name, preferred_competitors=parsed.competitors)
        emit("partial", section="linkedin", data=report.get("linkedin") or {}, client=client_name)
        _emit_progress(emit, set_step, "linkedin", "done", "LinkedIn ready")
    else:
        _emit_progress(emit, set_step, "linkedin", "done", "Skipped")

    # 7–9) Social
    _emit_progress(emit, set_step, "instagram", "running", "Instagram / YouTube / TikTok...")
    try:
        _run_script(
            [sys.executable, str(SCRIPT_SOCIAL), slug],
            FINAL_DIR,
            on_log,
            timeout_sec=600,
            env=script_env,
        )
    except Exception as e:
        emit("log", line=f"[SOCIAL] Partial failure: {e}")
    xlsx = find_metricas_xlsx(slug) or xlsx
    report = build_report_from_xlsx(xlsx, client_name=client_name, preferred_competitors=parsed.competitors)
    report = _attach_brand_dfs(report, slug)
    report = _attach_similarweb(report, slug)
    for sid in ("instagram", "youtube", "tiktok"):
        emit("partial", section=sid, data=report.get(sid) or {}, client=client_name)
        _emit_progress(emit, set_step, sid, "done", f"{sid} ready")

    html_url = ""
    if run_id:
        _save_report(
            run_id, report, parsed,
            site_only=site_only,
            include_google_ads=include_google_ads,
            include_linkedin=include_linkedin,
            force=force,
        )
        html_url = f"/runs/{run_id}/report.html"

    emit("done", report=report, html_url=html_url)
    return report
