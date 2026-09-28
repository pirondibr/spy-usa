# -*- coding: utf-8 -*-
"""Etapa Brand Search (DataForSEO) do Spy USA.

Uso:
    python "5g - brand dataforseo.py" semrush
    python "5g - brand dataforseo.py" semrush --out outputs/metricas/semrush/brand-dfs.json
"""

from __future__ import annotations

import json
import sys
from pathlib import Path

FINAL_DIR = Path(__file__).resolve().parent
VENDOR = FINAL_DIR / "vendor"
sys.path.insert(0, str(FINAL_DIR))
sys.path.insert(0, str(VENDOR))

from workspace_paths import setup_workspace, OUT_METRICAS, OUT_CONCORRENTES  # noqa: E402

setup_workspace()

from brand_dataforseo import brand_reports_batch, resolve_brand_keyword  # noqa: E402
from find_concorrentes import find_client_xlsx, read_briefing_from_xlsx  # noqa: E402
from seo_pipeline import scrape_site, brand_token_from_domain, normalize_domain  # noqa: E402


def _slugify(client: str) -> str:
    import re
    import unicodedata

    s = unicodedata.normalize("NFKD", client or "").encode("ascii", "ignore").decode("ascii")
    return re.sub(r"[^a-z0-9]+", "", s.lower()) or "cliente"


def _load_competitor_entities(slug: str, briefing: dict) -> list[dict]:
    folder = OUT_CONCORRENTES / slug
    entities: list[dict] = []
    client_url = briefing.get("url") or ""
    client_dom = normalize_domain(client_url)
    title = ""
    try:
        site = scrape_site(client_url) if client_url else {}
        title = (site.get("title") or "").strip()
    except Exception as e:
        print(f"[BRAND/DFS] scrape title fail: {e}", flush=True)

    client_name = (
        briefing.get("cliente")
        or briefing.get("client")
        or brand_token_from_domain(client_url)
        or slug
    )
    entities.append({
        "name": client_name,
        "company": client_name,
        "domain": client_dom,
        "url": client_url,
        "title": title,
        "is_client": True,
        "similaridade": "Cliente",
    })

    # Competitors from nacional xlsx if present
    xlsx = None
    if folder.exists():
        files = sorted(folder.glob(f"concorrentes-all-{slug}-nacional*.xlsx"), key=lambda p: p.stat().st_mtime, reverse=True)
        if not files:
            files = sorted(folder.glob(f"concorrentes-all-{slug}*.xlsx"), key=lambda p: p.stat().st_mtime, reverse=True)
        xlsx = files[0] if files else None

    if not xlsx:
        return entities

    try:
        import openpyxl
        wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
        if "Concorrentes Unificado" not in wb.sheetnames:
            wb.close()
            return entities
        ws = wb["Concorrentes Unificado"]
        rows = list(ws.iter_rows(values_only=True))
        wb.close()
        if not rows:
            return entities
        headers = [str(h or "").strip() for h in rows[0]]
        idx = {h: i for i, h in enumerate(headers)}

        def cell(row, name):
            i = idx.get(name)
            return row[i] if i is not None and i < len(row) else ""

        for row in rows[1:]:
            if not any(row):
                continue
            dom = normalize_domain(cell(row, "Dominio") or "")
            if not dom or dom == client_dom:
                continue
            sim = str(cell(row, "Similaridade") or "").strip()
            if sim.lower().replace("é", "e") not in ("alto", "medio"):
                continue
            title_c = str(cell(row, "Titulo") or "").strip()
            url = str(cell(row, "URL") or "") or f"https://{dom}/"
            name = title_c.split("|")[0].strip() if title_c else brand_token_from_domain(dom)
            entities.append({
                "name": name or dom,
                "company": name or dom,
                "domain": dom,
                "url": url,
                "title": title_c,
                "is_client": False,
                "similaridade": sim or "—",
            })
            if len(entities) >= 12:  # client + up to 11 competitors
                break
    except Exception as e:
        print(f"[BRAND/DFS] competitors load fail: {e}", flush=True)
    return entities


def main() -> None:
    if len(sys.argv) < 2:
        print('Uso: python "5g - brand dataforseo.py" <cliente>')
        sys.exit(1)
    client = sys.argv[1].strip()
    slug = _slugify(client)
    out_path = OUT_METRICAS / slug / f"brand-dataforseo-{slug}.json"
    for i, a in enumerate(sys.argv):
        if a == "--out" and i + 1 < len(sys.argv):
            out_path = Path(sys.argv[i + 1])

    briefing_xlsx = find_client_xlsx(client)
    briefing = read_briefing_from_xlsx(briefing_xlsx)
    entities = _load_competitor_entities(slug, briefing)
    print(f"[BRAND/DFS] entidades={len(entities)} slug={slug}", flush=True)
    for e in entities:
        meta = resolve_brand_keyword(
            e.get("url") or e.get("domain") or "",
            title=e.get("title") or "",
            company_hint=e.get("company") or e.get("name") or "",
        )
        print(
            f"  - {e.get('domain')} keyword='{meta['keyword']}' "
            f"title='{(e.get('title') or '')[:40]}'",
            flush=True,
        )

    rows = brand_reports_batch(entities)
    payload = {
        "slug": slug,
        "source": "dataforseo_google_ads",
        "market": "US",
        "location_code": 2840,
        "rows": rows,
        "client_keyword": next((r.get("brand_keyword") for r in rows if r.get("is_client")), ""),
        "latest_label": next((r.get("latest_label") for r in rows if r.get("is_client")), None),
    }
    out_path.parent.mkdir(parents=True, exist_ok=True)
    out_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"[OUT] {out_path}", flush=True)
    for r in rows[:8]:
        print(
            f"  {r.get('brand_keyword')}: {r.get('latest_label')}={r.get('latest_volume')} "
            f"1y={r.get('growth_1y_pct')}% 2y={r.get('growth_2y_pct')}% "
            f"3y={r.get('growth_3y_pct')}% 5y={r.get('growth_5y_pct')}%",
            flush=True,
        )


if __name__ == "__main__":
    main()
