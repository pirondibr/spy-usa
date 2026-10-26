# -*- coding: utf-8 -*-
"""Etapa Instagram Influencers (TinyFish Search+Fetch) — opcional no Spy USA."""

from __future__ import annotations

import json
import os
import sys
import time
from pathlib import Path
from urllib.parse import urlparse

FINAL_DIR = Path(__file__).resolve().parent
VENDOR = FINAL_DIR / "vendor"
if str(VENDOR) not in sys.path:
    sys.path.insert(0, str(VENDOR))
if str(FINAL_DIR) not in sys.path:
    sys.path.insert(0, str(FINAL_DIR))

from workspace_paths import OUT_METRICAS as METRICAS_DIR  # noqa: E402
from brand_dataforseo import influencer_search_query  # noqa: E402
from instagram_influencers import analyze_company  # noqa: E402


def _domain(url: str) -> str:
    try:
        host = urlparse(url if "://" in url else f"https://{url}").netloc.lower()
        return host[4:] if host.startswith("www.") else host
    except Exception:
        return ""


def _load_entities(slug: str) -> list[dict]:
    xlsx = METRICAS_DIR / slug / f"metricas-concorrentes-{slug}.xlsx"
    entities: list[dict] = []
    if xlsx.exists():
        try:
            from openpyxl import load_workbook

            wb = load_workbook(xlsx, read_only=True, data_only=True)
            if "Metricas Canais" in wb.sheetnames:
                ws = wb["Metricas Canais"]
                rows = list(ws.iter_rows(values_only=True))
                if rows:
                    headers = [str(h or "").strip().lower() for h in rows[0]]

                    def col(*names):
                        for n in names:
                            if n in headers:
                                return headers.index(n)
                        return None

                    i_dom = col("dominio", "domain")
                    i_url = col("url")
                    i_sim = col("similaridade")
                    i_ig = col("instagram url", "instagram")
                    for row in rows[1:]:
                        if not row:
                            continue
                        dom = str(row[i_dom] or "").strip() if i_dom is not None else ""
                        url = str(row[i_url] or "").strip() if i_url is not None else ""
                        sim = str(row[i_sim] or "").strip() if i_sim is not None else ""
                        ig = str(row[i_ig] or "").strip() if i_ig is not None else ""
                        if not dom and url:
                            dom = _domain(url)
                        if not dom:
                            continue
                        is_client = sim.lower() in ("cliente", "client", "you") or dom.replace(
                            ".", ""
                        ) == slug.replace("-", "").replace("_", "")
                        name = dom.split(".")[0]
                        entities.append({
                            "name": name,
                            "domain": dom,
                            "url": url or f"https://{dom}/",
                            "instagram_url": ig,
                            "is_client": is_client,
                            "similaridade": sim or ("Cliente" if is_client else "—"),
                        })
            wb.close()
        except Exception as e:
            print(f"[IG-INF] warn read xlsx: {e}", flush=True)

    if not entities:
        entities = [{
            "name": slug,
            "domain": "",
            "url": "",
            "instagram_url": "",
            "is_client": True,
            "similaridade": "Cliente",
        }]
    return entities


def main() -> None:
    if len(sys.argv) < 2:
        print('Uso: python "5j - instagram influencers.py" <cliente_slug> [query]')
        sys.exit(1)

    root = FINAL_DIR.parent
    for name in (".env", ".env.local"):
        path = root / name
        if not path.exists():
            continue
        for line in path.read_text(encoding="utf-8").splitlines():
            line = line.strip()
            if not line or line.startswith("#") or "=" not in line:
                continue
            k, _, v = line.partition("=")
            k, v = k.strip(), v.strip().strip('"').strip("'")
            if k and k not in os.environ:
                os.environ[k] = v

    slug = sys.argv[1].strip()
    os.environ["SPY_SLUG"] = slug
    override_query = sys.argv[2].strip() if len(sys.argv) >= 3 else ""
    market = (os.environ.get("SPY_MARKET") or "us").strip().lower()
    site_only = os.environ.get("SPY_SITE_ONLY", "").strip().lower() in ("1", "true", "yes")

    entities = _load_entities(slug)
    if site_only:
        client = next((e for e in entities if e.get("is_client")), entities[0])
        entities = [client]

    print(
        f"[IG-INF] slug={slug} entities={len(entities)} market={market} "
        f"site_only={site_only}",
        flush=True,
    )

    rows = []
    t0 = time.time()
    total_fetch_urls = 0

    for ent in entities:
        name = ent.get("name") or slug
        domain = ent.get("domain") or ""
        q = influencer_search_query(
            slug,
            name=name,
            domain=domain,
            url=str(ent.get("url") or ""),
            is_client=bool(ent.get("is_client")),
            override=override_query,
        )
        print(f"[IG-INF] -> {name} ({domain}) query={q!r}", flush=True)
        try:
            result = analyze_company(
                name,
                domain=domain,
                market=market,
                profile_url=ent.get("instagram_url") or None,
                query=q,
            )
        except Exception as e:
            print(f"[IG-INF] ERROR {name}: {e}", flush=True)
            result = {
                "ok": False,
                "company": name,
                "domain": domain,
                "error": str(e),
                "total_influencers": 0,
                "total_likes": 0,
                "top5": [],
                "usage": {},
                "elapsed_sec": 0,
                "profile_url": "",
                "handle": "",
            }
        usage = result.get("usage") or {}
        total_fetch_urls += int(usage.get("fetch_urls") or 0)
        print(
            f"[IG-INF]   handle=@{result.get('handle') or '—'} "
            f"influencers={result.get('total_influencers')} "
            f"posts={result.get('posts_first_page')} "
            f"likes={result.get('total_likes')} "
            f"t={result.get('elapsed_sec')}s",
            flush=True,
        )
        rows.append({
            "name": name,
            "domain": domain,
            "is_client": bool(ent.get("is_client")),
            "similaridade": ent.get("similaridade") or "—",
            "query": result.get("query") or q,
            "profile_url": result.get("profile_url") or "",
            "handle": result.get("handle") or "",
            "followers_label": result.get("followers_label"),
            "posts_first_page": result.get("posts_first_page") or 0,
            "own_posts": result.get("own_posts") or 0,
            "influencer_posts": result.get("influencer_posts") or 0,
            "total_influencers": result.get("total_influencers") or 0,
            "influencer_handles": result.get("influencer_handles") or [],
            "total_likes": result.get("total_likes") or 0,
            "own_likes": result.get("own_likes") or 0,
            "influencer_likes": result.get("influencer_likes") or 0,
            "top5": result.get("top5") or [],
            "top5_page": result.get("top5_page") or [],
            "top5_influencers": result.get("top5_influencers") or [],
            "usage": usage,
            "elapsed_sec": result.get("elapsed_sec"),
            "ok": result.get("ok"),
            "error": result.get("error"),
        })

    out = {
        "slug": slug,
        "market": market.upper(),
        "source": "tinyfish_search+fetch",
        "elapsed_sec": round(time.time() - t0, 2),
        "usage_total": {
            "estimated_usd": 0.0,
            "fetch_urls": total_fetch_urls,
            "note": "Search + Fetch download only (free within TinyFish daily Fetch allowance)",
        },
        "rows": rows,
    }
    out_dir = METRICAS_DIR / slug
    out_dir.mkdir(parents=True, exist_ok=True)
    out_path = out_dir / f"instagram-influencers-{slug}.json"
    out_path.write_text(json.dumps(out, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"[OUT] {out_path}", flush=True)
    print(
        f"[IG-INF] TOTAL {out['elapsed_sec']}s · fetch_urls={total_fetch_urls} · cost≈$0",
        flush=True,
    )


if __name__ == "__main__":
    main()
