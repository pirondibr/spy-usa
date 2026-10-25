# -*- coding: utf-8 -*-
"""Etapa TikTok Influencers (Apify keyword search + brand filter)."""

from __future__ import annotations

import json
import os
import sys
import time
from pathlib import Path
from urllib.parse import urlparse

FINAL_DIR = Path(__file__).resolve().parent
VENDOR = FINAL_DIR / "vendor"
if str(VENDOR) not in sys.path:
    sys.path.insert(0, str(VENDOR))
if str(FINAL_DIR) not in sys.path:
    sys.path.insert(0, str(FINAL_DIR))

from workspace_paths import OUT_METRICAS as METRICAS_DIR  # noqa: E402
from tiktok_influencers import analyze_company, official_hint_from_url  # noqa: E402


def _domain(url: str) -> str:
    try:
        host = urlparse(url if "://" in url else f"https://{url}").netloc.lower()
        return host[4:] if host.startswith("www.") else host
    except Exception:
        return ""


def _load_entities(slug: str) -> list[dict]:
    xlsx = METRICAS_DIR / slug / f"metricas-concorrentes-{slug}.xlsx"
    entities: list[dict] = []
    if xlsx.exists():
        try:
            from openpyxl import load_workbook

            wb = load_workbook(xlsx, read_only=True, data_only=True)
            if "Metricas Canais" in wb.sheetnames:
                ws = wb["Metricas Canais"]
                rows = list(ws.iter_rows(values_only=True))
                if rows:
                    headers = [str(h or "").strip().lower() for h in rows[0]]

                    def col(*names):
                        for n in names:
                            if n in headers:
                                return headers.index(n)
                        return None

                    i_dom = col("dominio", "domain")
                    i_url = col("url")
                    i_sim = col("similaridade")
                    i_tt = col("tiktok url", "tik tok url")
                    for row in rows[1:]:
                        if not row:
                            continue
                        dom = str(row[i_dom] or "").strip() if i_dom is not None else ""
                        url = str(row[i_url] or "").strip() if i_url is not None else ""
                        sim = str(row[i_sim] or "").strip() if i_sim is not None else ""
                        tt = str(row[i_tt] or "").strip() if i_tt is not None else ""
                        if not dom and url:
                            dom = _domain(url)
                        if not dom:
                            continue
                        is_client = sim.lower() in ("cliente", "client", "you") or dom.replace(
                            ".", ""
                        ) == slug.replace("-", "").replace("_", "")
                        name = dom.split(".")[0]
                        entities.append({
                            "name": name,
                            "domain": dom,
                            "url": url or f"https://{dom}/",
                            "tiktok_url": tt,
                            "is_client": is_client,
                            "similaridade": sim or ("Cliente" if is_client else "—"),
                        })
            wb.close()
        except Exception as e:
            print(f"[TT-INF] warn read xlsx: {e}", flush=True)

    if not entities:
        entities = [{
            "name": slug,
            "domain": "",
            "url": "",
            "tiktok_url": "",
            "is_client": True,
            "similaridade": "Cliente",
        }]
    return entities


def main() -> None:
    if len(sys.argv) < 2:
        print('Uso: python "5n - tiktok influencers.py" <cliente_slug> [query]')
        sys.exit(1)

    root = FINAL_DIR.parent
    for name in (".env", ".env.local"):
        path = root / name
        if not path.exists():
            continue
        for line in path.read_text(encoding="utf-8").splitlines():
            line = line.strip()
            if not line or line.startswith("#") or "=" not in line:
                continue
            k, _, v = line.partition("=")
            k, v = k.strip(), v.strip().strip('"').strip("'")
            if k and k not in os.environ:
                os.environ[k] = v

    slug = sys.argv[1].strip()
    override_query = sys.argv[2].strip() if len(sys.argv) >= 3 else ""
    market = (os.environ.get("SPY_MARKET") or "us").strip().lower()
    site_only = os.environ.get("SPY_SITE_ONLY", "").strip().lower() in ("1", "true", "yes")

    entities = _load_entities(slug)
    if site_only:
        client = next((e for e in entities if e.get("is_client")), entities[0])
        entities = [client]

    print(
        f"[TT-INF] slug={slug} entities={len(entities)} market={market} "
        f"site_only={site_only}",
        flush=True,
    )

    rows = []
    t0 = time.time()
    total_usd = 0.0

    for ent in entities:
        name = ent.get("name") or slug
        domain = ent.get("domain") or ""
        q = override_query if (override_query and ent.get("is_client")) else ""
        print(f"[TT-INF] -> {name} ({domain}) query={q or '(brand ai)'!r}", flush=True)
        try:
            result = analyze_company(
                name,
                domain=domain,
                market=market,
                query=q,
                official_hint=official_hint_from_url(str(ent.get("tiktok_url") or "")),
            )
        except Exception as e:
            print(f"[TT-INF] ERROR {name}: {e}", flush=True)
            result = {
                "ok": False,
                "company": name,
                "domain": domain,
                "error": str(e),
                "views_official": 0,
                "views_unofficial": 0,
                "top5": [],
                "top5_official": [],
                "top5_unofficial": [],
                "counts": {},
                "usage": {},
                "elapsed_sec": 0,
            }
        usage = result.get("usage") or {}
        total_usd += float(usage.get("estimated_usd") or 0)
        counts = result.get("counts") or {}
        src = (usage.get("source") or "none")
        print(
            f"[TT-INF]   confirmed={counts.get('confirmed')} "
            f"probable={counts.get('probable')} "
            f"unrelated={counts.get('unrelated')} "
            f"off={result.get('views_official')} "
            f"oth={result.get('views_unofficial')} "
            f"src={src} cached={usage.get('cached')} "
            f"t={result.get('elapsed_sec')}s",
            flush=True,
        )
        rows.append({
            "name": name,
            "domain": domain,
            "is_client": bool(ent.get("is_client")),
            "similaridade": ent.get("similaridade") or "—",
            "query": result.get("query") or q,
            "search_url": result.get("search_url") or "",
            "views_official": result.get("views_official") or 0,
            "views_unofficial": result.get("views_unofficial") or 0,
            "views_confirmed": result.get("views_confirmed") or 0,
            "views_probable": result.get("views_probable") or 0,
            "about_count": counts.get("confirmed") or counts.get("about_product") or 0,
            "page_count": counts.get("search_page") or 0,
            "confirmed_count": counts.get("confirmed") or 0,
            "probable_count": counts.get("probable") or 0,
            "unrelated_count": counts.get("unrelated") or 0,
            "top5": result.get("top5") or [],
            "top5_official": result.get("top5_official") or [],
            "top5_unofficial": result.get("top5_unofficial") or [],
            "top5_probable": result.get("top5_probable") or [],
            "videos_confirmed": result.get("videos_confirmed") or [],
            "videos_probable": result.get("videos_probable") or [],
            "videos_unrelated": result.get("videos_unrelated") or [],
            "videos_fetched": result.get("videos_fetched") or [],
            "official_channel": result.get("official_channel"),
            "brand": result.get("brand") or {},
            "usage": usage,
            "elapsed_sec": result.get("elapsed_sec"),
            "ok": result.get("ok"),
            "error": result.get("error"),
        })

    src = next(
        (
            str((r.get("usage") or {}).get("source") or "")
            for r in rows
            if (r.get("usage") or {}).get("source")
        ),
        "apify_tiktok",
    )
    out = {
        "slug": slug,
        "market": market.upper(),
        "source": src or "apify_tiktok",
        "elapsed_sec": round(time.time() - t0, 2),
        "usage_total": {
            "estimated_usd": round(total_usd, 4),
            "search_cost_usd": round(total_usd, 4),
            "source": src or "apify_tiktok",
            "note": "Apify apidojo/tiktok-scraper. Brand filter hides unrelated. ~$0.0075 / 25 posts.",
        },
        "brand_filter": "v2_neg_strong_brand",
        "rows": rows,
    }
    out_dir = METRICAS_DIR / slug
    out_dir.mkdir(parents=True, exist_ok=True)
    out_path = out_dir / f"tiktok-influencers-{slug}.json"
    out_path.write_text(json.dumps(out, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"[OUT] {out_path}", flush=True)
    print(
        f"[TT-INF] TOTAL {out['elapsed_sec']}s · "
        f"src={out.get('source')} · ~${total_usd:.4f}",
        flush=True,
    )


if __name__ == "__main__":
    main()
