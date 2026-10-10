# -*- coding: utf-8 -*-
"""Etapa Twitter/X profile (TinyFish Search + ScrapingBee followers)."""

from __future__ import annotations

import json
import os
import sys
import time
from pathlib import Path
from urllib.parse import urlparse

FINAL_DIR = Path(__file__).resolve().parent
VENDOR = FINAL_DIR / "vendor"
for p in (VENDOR, FINAL_DIR):
    if str(p) not in sys.path:
        sys.path.insert(0, str(p))

from workspace_paths import OUT_METRICAS as METRICAS_DIR  # noqa: E402
from twitter_profile import analyze_company  # noqa: E402


def _domain(url: str) -> str:
    try:
        host = urlparse(url if "://" in url else f"https://{url}").netloc.lower()
        return host[4:] if host.startswith("www.") else host
    except Exception:
        return ""


def _load_entities(slug: str) -> list[dict]:
    xlsx = METRICAS_DIR / slug / f"metricas-concorrentes-{slug}.xlsx"
    entities: list[dict] = []
    if xlsx.exists():
        try:
            from openpyxl import load_workbook

            wb = load_workbook(xlsx, read_only=True, data_only=True)
            if "Metricas Canais" in wb.sheetnames:
                ws = wb["Metricas Canais"]
                rows = list(ws.iter_rows(values_only=True))
                if rows:
                    headers = [str(h or "").strip().lower() for h in rows[0]]

                    def col(*names):
                        for n in names:
                            if n in headers:
                                return headers.index(n)
                        return None

                    i_dom, i_url, i_sim = col("dominio", "domain"), col("url"), col("similaridade")
                    for row in rows[1:]:
                        if not row:
                            continue
                        dom = str(row[i_dom] or "").strip() if i_dom is not None else ""
                        url = str(row[i_url] or "").strip() if i_url is not None else ""
                        sim = str(row[i_sim] or "").strip() if i_sim is not None else ""
                        if not dom and url:
                            dom = _domain(url)
                        if not dom:
                            continue
                        is_client = sim.lower() in ("cliente", "client", "you") or dom.replace(
                            ".", ""
                        ) == slug.replace("-", "").replace("_", "")
                        entities.append({
                            "name": dom.split(".")[0],
                            "domain": dom,
                            "url": url or f"https://{dom}/",
                            "is_client": is_client,
                            "similaridade": sim or ("Cliente" if is_client else "—"),
                        })
            wb.close()
        except Exception as e:
            print(f"[TW] warn read xlsx: {e}", flush=True)
    if not entities:
        entities = [{"name": slug, "domain": "", "url": "", "is_client": True, "similaridade": "Cliente"}]
    return entities


def main() -> None:
    if len(sys.argv) < 2:
        print('Uso: python "5l - twitter profile.py" <cliente_slug> [query]')
        sys.exit(1)
    root = FINAL_DIR.parent
    for name in (".env", ".env.local"):
        path = root / name
        if not path.exists():
            continue
        for line in path.read_text(encoding="utf-8").splitlines():
            line = line.strip()
            if not line or line.startswith("#") or "=" not in line:
                continue
            k, _, v = line.partition("=")
            k, v = k.strip(), v.strip().strip('"').strip("'")
            if k and k not in os.environ:
                os.environ[k] = v

    slug = sys.argv[1].strip()
    override_query = sys.argv[2].strip() if len(sys.argv) >= 3 else ""
    market = (os.environ.get("SPY_MARKET") or "us").strip().lower()
    site_only = os.environ.get("SPY_SITE_ONLY", "").strip().lower() in ("1", "true", "yes")
    entities = _load_entities(slug)
    if site_only:
        entities = [next((e for e in entities if e.get("is_client")), entities[0])]

    print(f"[TW] slug={slug} entities={len(entities)} market={market}", flush=True)
    rows, total_credits = [], 0
    t0 = time.time()
    for ent in entities:
        name = ent.get("name") or slug
        domain = ent.get("domain") or ""
        q = override_query if (override_query and ent.get("is_client")) else name.replace("-", " ").strip()
        print(f"[TW] -> {name} ({domain}) query={q!r}", flush=True)
        try:
            result = analyze_company(name, domain=domain, market=market, query=q)
        except Exception as e:
            print(f"[TW] ERROR {name}: {e}", flush=True)
            result = {"ok": False, "error": str(e), "followers": None, "usage": {}}
        usage = result.get("usage") or {}
        total_credits += int(usage.get("scrapingbee_credits") or 0)
        print(
            f"[TW]   @{result.get('handle') or '—'} "
            f"followers={result.get('followers_text') or result.get('followers')} "
            f"src={result.get('followers_source')} t={result.get('elapsed_sec')}s",
            flush=True,
        )
        rows.append({
            "name": name,
            "domain": domain,
            "is_client": bool(ent.get("is_client")),
            "similaridade": ent.get("similaridade") or "—",
            "query": result.get("query") or q,
            "search_url": result.get("search_url") or "",
            "profile_url": result.get("profile_url") or "",
            "handle": result.get("handle") or "",
            "followers": result.get("followers"),
            "followers_text": result.get("followers_text") or "",
            "followers_source": result.get("followers_source"),
            "title": result.get("title") or "",
            "usage": usage,
            "elapsed_sec": result.get("elapsed_sec"),
            "ok": result.get("ok"),
            "error": result.get("error"),
        })

    out = {
        "slug": slug,
        "market": market.upper(),
        "source": "tinyfish_search + scrapingbee_profile",
        "elapsed_sec": round(time.time() - t0, 2),
        "usage_total": {"scrapingbee_credits": total_credits},
        "rows": rows,
    }
    out_dir = METRICAS_DIR / slug
    out_dir.mkdir(parents=True, exist_ok=True)
    out_path = out_dir / f"twitter-profile-{slug}.json"
    out_path.write_text(json.dumps(out, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"[OUT] {out_path}", flush=True)
    print(f"[TW] TOTAL {out['elapsed_sec']}s · sb_credits={total_credits}", flush=True)


if __name__ == "__main__":
    main()
